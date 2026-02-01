
import os
import gzip
from typing import Dict, Optional, Iterable, Tuple

from openpyxl import load_workbook

# ========= CONFIG =========
INPUT_XLSX = r"C:\Users\moise\OneDrive - Universidade do Algarve\Biotec 2\GProtein\Callorhinchus milii\Callorhinchus.milli_FamilyB1_selected_and_fasta.xlsx"
OUTPUT_XLSX = r"C:\Users\moise\OneDrive - Universidade do Algarve\Biotec 2\GProtein\Callorhinchus milii\Callorhinchus.milli_FamilyB1_with_ExonCount.xlsx"

SHEETS = ["CALCR", "CRHR", "PTHR", "GCGR", "VIP SCRT"]  # ou None para todas
COL_ACCESSION = "A"   # XP_...
COL_EXONCOUNT = "X"   # saída

# Caminho para o GFF3 do assembly (do zip do datasets/FTP)
GFF_PATH = r"C:\Users\moise\OneDrive - Universidade do Algarve\Biotec 2\GProtein\Callorhinchus milii\GCF_018977255.1_IMCB_Cmil_1.0_genomic.gff.gz"  # <-- AJUSTA para o teu ficheiro .gff ou .gff.gz

# ========= GFF PARSING =========
def open_text_maybe_gz(path: str) -> Iterable[str]:
    if path.lower().endswith(".gz"):
        with gzip.open(path, "rt", encoding="utf-8", errors="replace") as f:
            for line in f:
                yield line
    else:
        with open(path, "rt", encoding="utf-8", errors="replace") as f:
            for line in f:
                yield line

def parse_gff_attributes(attr: str) -> Dict[str, str]:
    """
    Parse básico de atributos GFF3: key=value;key2=value2
    Pode haver múltiplos valores separados por vírgulas (guardamos string crua).
    """
    out: Dict[str, str] = {}
    for part in attr.strip().split(";"):
        if not part:
            continue
        if "=" not in part:
            continue
        k, v = part.split("=", 1)
        out[k] = v
    return out

def first_parent(parent_field: str) -> Optional[str]:
    if not parent_field:
        return None
    # Parent pode ter múltiplos IDs separados por vírgula
    return parent_field.split(",")[0].strip() or None

def build_maps_from_gff(gff_path: str) -> Tuple[Dict[str, int], Dict[str, str]]:
    """
    Devolve:
      exon_count_by_tx: transcript_id -> nº exons
      tx_by_protein:   protein_id (XP_...) -> transcript_id (Parent)
    """
    exon_count_by_tx: Dict[str, int] = {}
    tx_by_protein: Dict[str, str] = {}

    for line in open_text_maybe_gz(gff_path):
        if not line or line.startswith("#"):
            continue
        cols = line.rstrip("\n").split("\t")
        if len(cols) != 9:
            continue

        seqid, source, ftype, start, end, score, strand, phase, attrs = cols
        a = parse_gff_attributes(attrs)

        if ftype == "exon":
            tx = first_parent(a.get("Parent", ""))
            if tx:
                exon_count_by_tx[tx] = exon_count_by_tx.get(tx, 0) + 1

        elif ftype == "CDS":
            # RefSeq GFF3 costuma ter protein_id=XP_...
            prot = a.get("protein_id") or a.get("proteinId") or a.get("protein", "")
            prot = prot.split(",")[0].strip() if prot else ""
            tx = first_parent(a.get("Parent", ""))

            # fallback: às vezes vem em Dbxref=Genbank:XP_...
            if (not prot) and ("Dbxref" in a):
                for item in a["Dbxref"].split(","):
                    item = item.strip()
                    if item.startswith("Genbank:") and item.split("Genbank:", 1)[1].startswith(("XP_", "NP_", "YP_")):
                        prot = item.split("Genbank:", 1)[1]
                        break

            if prot and tx:
                # ficar com o primeiro mapeamento (deve ser 1:1)
                tx_by_protein.setdefault(prot, tx)

    return exon_count_by_tx, tx_by_protein

def get_exon_count_for_protein(prot_acc: str, exon_count_by_tx: Dict[str, int], tx_by_protein: Dict[str, str]) -> Optional[int]:
    tx = tx_by_protein.get(prot_acc)
    if not tx:
        return None
    return exon_count_by_tx.get(tx)

# ========= MAIN =========
def main():
    if not os.path.exists(GFF_PATH):
        raise SystemExit(f"Não encontrei o GFF em: {GFF_PATH}")

    print("A indexar GFF3 (isto pode demorar, dependendo do tamanho)...")
    exon_count_by_tx, tx_by_protein = build_maps_from_gff(GFF_PATH)
    print(f"  transcripts com exons: {len(exon_count_by_tx)}")
    print(f"  proteínas mapeadas:    {len(tx_by_protein)}")

    wb = load_workbook(INPUT_XLSX)
    sheetnames = SHEETS if SHEETS is not None else wb.sheetnames

    for sh in sheetnames:
        if sh not in wb.sheetnames:
            print(f"[AVISO] Folha '{sh}' não existe. A saltar.")
            continue

        ws = wb[sh]
        max_row = ws.max_row
        print(f"\n==> {sh} | linhas: {max_row}")

        for row in range(2, max_row + 1):  # linha 1 = cabeçalho
            acc_val = ws[f"{COL_ACCESSION}{row}"].value
            if acc_val is None:
                continue

            prot = str(acc_val).strip()
            if not prot:
                continue

            exon_count = get_exon_count_for_protein(prot, exon_count_by_tx, tx_by_protein)
            ws[f"{COL_EXONCOUNT}{row}"].value = exon_count  # pode ficar None se não mapear

            if row % 50 == 0:
                print(f"  ... {row}/{max_row}")

    wb.save(OUTPUT_XLSX)
    print(f"\n[OK] Guardado: {OUTPUT_XLSX}")

if __name__ == "__main__":
    main()
