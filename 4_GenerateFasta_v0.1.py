from openpyxl import load_workbook

# ========= CONFIG =========
INPUT_XLSX = r"C:\Users\moise\OneDrive - Universidade do Algarve\Biotec 2\GProtein\Latimeria chalumnae\Latimeria.chalumnae_FamilyB1_with_ExonCount.xlsx"
OUTPUT_FASTA = r"C:\Users\moise\OneDrive - Universidade do Algarve\Biotec 2\GProtein\Latimeria chalumnae\Latimeria_chalumnae_GPCR.fasta"

COL_FASTA = "Z"       # coluna onde está o texto FASTA completo
HEADER_ROW = 1        # linha de cabeçalho (vai ser ignorada)


def append_sheet_to_headers(fasta_text: str, sheet_tag: str) -> str:
    """
    Recebe um FASTA (texto) e acrescenta _{sheet_tag} (ex: _CALCR) a todas as linhas de header.
    Mantém o resto exatamente igual.
    """
    out_lines = []
    for line in fasta_text.splitlines():
        line = line.rstrip("\r\n")
        if line.startswith(">"):
            # Evita duplicar caso já tenha sido processado antes
            if not line.endswith(f"_{sheet_tag}"):
                line = f"{line}_{sheet_tag}"
        out_lines.append(line)
    return "\n".join(out_lines).strip()


def main():
    wb = load_workbook(INPUT_XLSX)

    with open(OUTPUT_FASTA, "w", encoding="utf-8") as fout:
        # Todas as folhas menos a primeira
        for sheet_name in wb.sheetnames[1:]:
            ws = wb[sheet_name]
            sheet_tag = sheet_name.upper()

            for row in range(HEADER_ROW + 1, ws.max_row + 1):  # ignora cabeçalho
                cell_val = ws[f"{COL_FASTA}{row}"].value
                if not cell_val:
                    continue

                fasta_text = str(cell_val).strip()
                if not fasta_text:
                    continue

                # Se por algum motivo não começar por '>', ainda escreve, mas não mexe
                if ">" in fasta_text:
                    fasta_text = append_sheet_to_headers(fasta_text, sheet_tag)

                # Garante separação entre registos
                fout.write(fasta_text + "\n")
                if not fasta_text.endswith("\n"):
                    fout.write("\n")

    print(f"[OK] FASTA criado: {OUTPUT_FASTA}")


if __name__ == "__main__":
    main()
